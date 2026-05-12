# -*- coding: utf-8 -*-
"""
生成购销合同 Word 模板（A4 横向，严格匹配 carote 截图样式）。

运行：python build_template.py
输出：template/contract_template.docx

模板内置的 docxtpl 占位符：
  - {{ 甲方 }}、{{ 乙方 }}、{{ 合同号码 }}、{{ 签约日期 }}
  - 表格行循环：{%tr for item in items %} ... {%tr endfor %}
    每行变量：item.采购订单号 / 采购行号 / 物料编码 / 物料名称 / 交期 / 数量 / 金额
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "template"
TEMPLATE_PATH = TEMPLATE_DIR / "contract_template.docx"
ASSETS_DIR = BASE_DIR / "assets"

# 优先用 jpg；同时兼容 png
_LOGO_CANDIDATES = [
    ASSETS_DIR / "carote_logo.jpg",
    ASSETS_DIR / "carote_logo.png",
]


def _logo_path() -> Path | None:
    for p in _LOGO_CANDIDATES:
        if p.is_file():
            return p
    return None


def _set_run(
    run,
    *,
    font: str = "微软雅黑",
    size_pt: float = 9,
    bold: bool = False,
    color: tuple[int, int, int] = (0, 0, 0),
):
    """统一设置 run 字体/大小/颜色（含中文东亚字体）。"""
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size_pt)
    run.bold = bold
    r, g, b = color
    run.font.color.rgb = RGBColor(r, g, b)


def _set_cell_borders(cell, color: str = "B5B5B5", size: str = "4"):
    """单元格四边边框：颜色 6 位 hex；size 单位 1/8 pt。"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        e = tcBorders.find(qn(f"w:{edge}"))
        if e is None:
            e = OxmlElement(f"w:{edge}")
            tcBorders.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:color"), color)


def _clear_cell_borders(cell):
    """移除单元格全部边框（用于无边框的标题/签署区表格）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        e = tcBorders.find(qn(f"w:{edge}"))
        if e is None:
            e = OxmlElement(f"w:{edge}")
            tcBorders.append(e)
        e.set(qn("w:val"), "nil")


def _set_cell_shading(cell, fill_hex: str = "F5F5F5"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def _set_col_width(table, widths_mm: list[int]):
    """设置每列固定宽度（同时设置 tblLayout=fixed）。"""
    tblPr = table._element.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._element.find(qn("w:tblGrid"))
    if grid is not None:
        for el in list(grid):
            grid.remove(el)
        for w in widths_mm:
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(int(Mm(w).twips)))
            grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_mm):
                cell.width = Mm(widths_mm[i])


def _add_term_paragraph(doc, idx: int, text: str, *, size_pt: float = 8.5):
    """添加一条以「N、」开头的小字条款段落。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r0 = p.add_run(f"{idx}、")
    _set_run(r0, size_pt=size_pt, bold=True)
    r1 = p.add_run(text)
    _set_run(r1, size_pt=size_pt)
    return p


# ----- 条款文本（请按贵司法务最终版替换；当前文本贴近截图视觉效果，仅供排版参考）-----
PRODUCT_TERMS: list[str] = [
    "材料要求：保证全部按照欧 EN801 标准，如同材料成份、涂层性能、磨损及外观等给消费者满意的体感不会变，责任及后果由乙方负责。",
    "模具进度及参数：由乙方承担模具的检测费及在正常使用情况的修补费等，乙方在收到甲方设备情况后，结合甲方设备情况，结合甲方需求情况后立即按需开模生产。",
    "货款付清：先开订单。",
    "验收方法：乙方为生产商提供产品制造时，乙方必须先填写出货前两周以申请甲方验收。如检验合格，则甲方签字同意乙方出货。",
    "质量要求：产品的外观无凹陷、无气泡、无裂痕、无变形、表面涂层均匀，提供电源 FCB 反复试验，如第一次未合格，则第二次返工生产，再不合格视为乙方违约。甲方在收到乙方提供的产品后检验入库时，检验合格后视为乙方交付义务完成；但实收合格不视为对产品质量的认可，如乙方产品品质出现问题视乎严重情况，乙方需返还相应金额损失乙方支付的尾款，可在乙方支付的款项中扣除。支付款不足以扣中甲方损失，乙方应继续追加补偿。",
    "交、运货方式：乙方工厂卡板验收装运，乙方负责装柜及发货前的报关运货等出口费用。甲方负责相应税收支出。",
    "货款支付：乙方在工厂生产经验过程中，乙方负责相应税收支出。乙方必须在生产 15 天内开具增值税发票给甲方。如果发生改收的款项需要乙方对款项的合法的补偿要求，因故意调整而引起的合同金额额变动金额超过销售标准，将本合同项下并货款顶进改成贷货支出，贷款按现金 90% 来打算。",
    "违约责任：乙方对于自身因故而延误生效之日，证明 7 日内向乙方告知，超过 0.3% 每日罚款，扩展 10 天以上，甲方有权回敕合同，一次性返货。如发生法投乙方拘束 250 美元。",
    "进货付款：约定 100% 完工后开发合并货款付款。",
]

OTHER_TERMS: list[str] = [
    "本合同未尽事宜，甲乙双方协商解决；如协商不成则提交甲方所在地法院核诉。",
    "本合同一式二份，双方各一份。",
    "乙方不得使用甲方的品牌或其他任何公司或个人对外提供相同或类似大宗产品，也不得向甲方提供的甲方客户的销售、设计图纸用作乙方的商业营销，一经发现，违者处以十万元人民币。",
    "乙方不得就此合同制卡 / 制盒 / 制贴制作产品成品内外排者于甲方以外的任何第三人，违不得用于任何用途的展示。如有违反此条款，则乙方按此合同总金额的 40% 向甲方支付违约金。",
]


def main():
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # ===== 页面：A4 横向 =====
    sec = doc.sections[0]
    sec.orientation = WD_ORIENTATION.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.left_margin = Mm(14)
    sec.right_margin = Mm(14)
    sec.top_margin = Mm(12)
    sec.bottom_margin = Mm(12)

    # 默认正文样式
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(9)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)

    # ===== 顶部三栏：logo | 标题 | 合同号 / 签约日期 =====
    # 总宽度 269mm（297 - 14*2）
    head = doc.add_table(rows=1, cols=3)
    head.autofit = False
    _set_col_width(head, [80, 109, 80])
    for c in head.rows[0].cells:
        _clear_cell_borders(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # logo
    p_logo = head.cell(0, 0).paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo = _logo_path()
    if logo is not None:
        p_logo.add_run().add_picture(str(logo), width=Mm(40))
    else:
        r = p_logo.add_run("carote")
        _set_run(r, font="Arial Black", size_pt=30, bold=True, color=(35, 35, 35))

    # 标题
    p_title = head.cell(0, 1).paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("购 销 合 同")
    _set_run(r, size_pt=28, bold=True, color=(0, 0, 0))

    # 合同号 / 签约日期
    cell_meta = head.cell(0, 2)
    p1 = cell_meta.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("合同号码：")
    _set_run(r1, size_pt=9, color=(120, 120, 120))
    r2 = p1.add_run("{{ 合同号码 }}")
    _set_run(r2, size_pt=9, color=(0, 0, 0))
    p2 = cell_meta.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    r3 = p2.add_run("签约日期：")
    _set_run(r3, size_pt=9, color=(120, 120, 120))
    r4 = p2.add_run("{{ 签约日期 }}")
    _set_run(r4, size_pt=9, color=(0, 0, 0))

    # ===== 甲方 / 乙方 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r1 = p.add_run("甲方：")
    _set_run(r1, size_pt=10, color=(120, 120, 120))
    r2 = p.add_run("{{ 甲方 }}")
    _set_run(r2, size_pt=10, bold=True)

    p = doc.add_paragraph()
    r1 = p.add_run("乙方：")
    _set_run(r1, size_pt=10, color=(120, 120, 120))
    r2 = p.add_run("{{ 乙方 }}")
    _set_run(r2, size_pt=10, bold=True)

    # ===== 引言 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("经甲乙双方友好协商，对下列产品名称、规格、单位、单价等达成如下协议内容，甲乙双方共同遵照履行：")
    _set_run(r, size_pt=9)

    # ===== 一、产品明细 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("一、产品明细：")
    _set_run(r, size_pt=10, bold=True)

    # ===== 商品表（带 docxtpl 行循环）=====
    # 采用四行模板：表头 / {%tr for%} / 数据 / {%tr endfor%}
    # for/endfor 行在渲染时被 docxtpl 整行删除，仅保留按 items 循环出的数据行。
    headers = ["采购订单号", "采购行号", "物料编码", "物料名称", "交期", "数量", "金额"]
    field_keys = ["采购订单号", "采购行号", "物料编码", "物料名称", "交期", "数量", "金额"]
    # 总宽 269mm
    col_widths = [35, 18, 24, 80, 24, 18, 70]

    table = doc.add_table(rows=4, cols=len(headers))
    table.autofit = False
    _set_col_width(table, col_widths)

    # 表头
    for i, txt in enumerate(headers):
        cell = table.cell(0, i)
        _set_cell_shading(cell, "F0F0F0")
        _set_cell_borders(cell, "B5B5B5", "4")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txt)
        _set_run(r, size_pt=9, bold=True)

    # for 行：仅放 {%tr for item in items %} 指令，整行将被 docxtpl 删除
    cell_for = table.cell(1, 0)
    pf = cell_for.paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.paragraph_format.space_after = Pt(0)
    rf = pf.add_run("{%tr for item in items %}")
    _set_run(rf, size_pt=1)

    # 数据行：每列填 {{ item.字段 }}
    row_data = table.rows[2]
    for i, key in enumerate(field_keys):
        cell = row_data.cells[i]
        _set_cell_borders(cell, "D5D5D5", "4")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("{{ item." + key + " }}")
        _set_run(r, size_pt=9)

    # endfor 行
    cell_end = table.cell(3, 0)
    pe = cell_end.paragraphs[0]
    pe.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pe.paragraph_format.space_after = Pt(0)
    re_ = pe.add_run("{%tr endfor %}")
    _set_run(re_, size_pt=1)

    # ===== 产品要求条款（1~9）=====
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(0)

    for i, t in enumerate(PRODUCT_TERMS, 1):
        _add_term_paragraph(doc, i, t, size_pt=8.5)

    # ===== 二、其他事项 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("二、其他事项：")
    _set_run(r, size_pt=10, bold=True)

    for i, t in enumerate(OTHER_TERMS, 1):
        _add_term_paragraph(doc, i, t, size_pt=8.5)

    # ===== 底部签署区 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)

    sig = doc.add_table(rows=4, cols=2)
    sig.autofit = False
    _set_col_width(sig, [134, 135])
    for r_ in sig.rows:
        for c in r_.cells:
            _clear_cell_borders(c)
            c.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    def _put(cell, label: str, value: str = ""):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label)
        _set_run(r1, size_pt=10, color=(80, 80, 80))
        if value:
            r2 = p.add_run(value)
            _set_run(r2, size_pt=10, bold=True)

    _put(sig.cell(0, 0), "甲方：", "{{ 甲方 }}")
    _put(sig.cell(0, 1), "乙方：", "{{ 乙方 }}")
    _put(sig.cell(1, 0), "电话：")
    _put(sig.cell(1, 1), "电话：")
    _put(sig.cell(2, 0), "法人代表（代理人）：")
    _put(sig.cell(2, 1), "法人代表（代理人）：")
    _put(sig.cell(3, 0), "")
    _put(sig.cell(3, 1), "")

    doc.save(str(TEMPLATE_PATH))
    print(f"模板已生成：{TEMPLATE_PATH}")


if __name__ == "__main__":
    main()
